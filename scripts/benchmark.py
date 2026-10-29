"""性能基线脚本 — 合成文档的 转换 / 审查 / 报告 耗时基准（2026-10-29）。

为什么需要：项目此前只有 5~6 页样例，没有任何大文档耗时数据，无法判断
"100 页审查是否可接受"或优化是否有效。本脚本生成确定性合成文档，分段计时
并输出可比较的表格（可选 JSON），用于优化前后对比与容量规划。

用法:
    python scripts/benchmark.py                       # docx+pptx, 10/50/100 页
    python scripts/benchmark.py --sizes 20 --format docx
    python scripts/benchmark.py --json logs/bench.json

注意：全部在本地临时目录生成并清理；LanguageTool 不可用时自动降级
（其探测耗时计入 audit 段，属于真实运行成本）。
"""

import argparse
import json
import random
import sys
import tempfile
import time
from pathlib import Path

# 直接运行脚本时保证 src/ 可导入（与 cli.py 同样处理）
PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

DEFAULT_SIZES = [10, 50, 100]
RULE_SUFFIXES = ("。", "；", "。", "")

_SENTENCE_POOL = [
    "本页说明关键工艺参数与量测条件，结论与上一阶段保持一致",
    "设备型号与批次信息见附表，数据来源于在线量测系统 MES",
    "良率波动主要来自刻蚀均匀性，已在下一轮 DOE 中安排验证",
    "所有数值均按 95% 置信区间统计，样本量不少于 30 片晶圆",
    "本节结论：工艺窗口满足量产要求，建议进入下一阶段验证",
    "注意事项：参数变更需同步更新控制计划与作业指导书",
]


def _body_paragraphs(rng: random.Random, count: int) -> list[str]:
    """生成 count 条中文正文句（确定性）。"""
    sentences = []
    for _ in range(count):
        sentences.append(rng.choice(_SENTENCE_POOL) + rng.choice(RULE_SUFFIXES))
    return sentences


def generate_pptx(path: Path, pages: int, seed: int = 42) -> None:
    """生成 pages 页 PPTX（每页标题 + 4 段正文 + 每 5 页一个表格）。"""
    from pptx import Presentation
    from pptx.util import Pt

    rng = random.Random(seed)
    prs = Presentation()
    for idx in range(pages):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"第 {idx + 1} 页 工艺评审"
        body = slide.placeholders[1].text_frame
        body.text = _body_paragraphs(rng, 1)[0]
        for sentence in _body_paragraphs(rng, 3):
            body.add_paragraph().text = sentence
        if idx % 5 == 0:
            table = slide.shapes.add_table(2, 2, Pt(50), Pt(420), Pt(400), Pt(60)).table
            table.cell(0, 0).text = "参数"
            table.cell(0, 1).text = "数值"
            table.cell(1, 0).text = "厚度"
            table.cell(1, 1).text = "12.5 nm"
    prs.save(str(path))


def generate_docx(path: Path, pages: int, seed: int = 42) -> None:
    """生成 pages 个章节的 DOCX（每章 H1 标题 + 4 段正文，按 H1 分页）。"""
    from docx import Document as DocxDocument

    rng = random.Random(seed)
    doc = DocxDocument()
    for idx in range(pages):
        doc.add_heading(f"第 {idx + 1} 章 工艺评审", level=1)
        for sentence in _body_paragraphs(rng, 4):
            doc.add_paragraph(sentence)
    doc.save(str(path))


def measure(path: Path, *, with_html: bool = True) -> dict:
    """测量单文件 转换 / 审查 / 报告 耗时（秒）与发现数。"""
    from src.engines.pipeline import build_auditors, find_converter, run_auditors
    from src.reporters.html_reporter import generate_html_report

    converter = find_converter(str(path))
    if converter is None:
        raise ValueError(f"无可用转换器: {path}")

    started = time.perf_counter()
    doc = converter.convert(str(path))
    convert_s = time.perf_counter() - started

    started = time.perf_counter()
    auditors = build_auditors(
        str(PROJECT_ROOT / "rules.md"),
        str(PROJECT_ROOT / "glossary"),
        str(PROJECT_ROOT / "vocab"),
    )
    findings = run_auditors(doc, auditors)
    audit_s = time.perf_counter() - started

    html_s = 0.0
    if with_html:
        started = time.perf_counter()
        generate_html_report(doc, findings)
        html_s = time.perf_counter() - started

    return {
        "file": path.name,
        "pages": len(doc.pages),
        "bytes": path.stat().st_size,
        "convert_s": round(convert_s, 3),
        "audit_s": round(audit_s, 3),
        "html_s": round(html_s, 3),
        "findings": len(findings),
    }


def run_benchmark(
    sizes: list[int] | None = None,
    formats: tuple[str, ...] = ("docx", "pptx"),
    seed: int = 42,
    with_html: bool = True,
    warmup: bool = True,
) -> list[dict]:
    """对每个 (格式, 页数) 组合生成文档并测量，返回结果列表。"""
    sizes = sizes or DEFAULT_SIZES
    results: list[dict] = []
    with tempfile.TemporaryDirectory(prefix="docaudit-bench-") as tmp:
        tmp_dir = Path(tmp)
        if warmup:
            # 预热：LanguageTool 探活缓存/拼写词典/导入开销只应计入"冷启动"，
            # 不应污染各页数的对比（实测首行审计 1.2s，后续 0.1s 量级）。
            warm_path = tmp_dir / "warmup.docx"
            generate_docx(warm_path, 1, seed)
            measure(warm_path, with_html=False)
        for fmt in formats:
            for pages in sizes:
                target = tmp_dir / f"bench_{pages}.{fmt}"
                if fmt == "pptx":
                    generate_pptx(target, pages, seed)
                elif fmt == "docx":
                    generate_docx(target, pages, seed)
                else:
                    raise ValueError(f"不支持格式: {fmt}")
                row = measure(target, with_html=with_html)
                row["format"] = fmt
                row["pages_requested"] = pages
                results.append(row)
    return results


def _print_table(results: list[dict]) -> None:
    header = (
        f"{'格式':<6}{'页数':>6}{'转换(s)':>10}{'审查(s)':>10}{'HTML(ms)':>10}"
        f"{'合计(s)':>10}{'发现':>8}{'大小(KB)':>10}"
    )
    print(header)
    print("-" * len(header))
    for row in results:
        total = row["convert_s"] + row["audit_s"] + row["html_s"]
        print(
            f"{row['format']:<6}{row['pages']:>6}{row['convert_s']:>10.2f}"
            f"{row['audit_s']:>10.2f}{row['html_s'] * 1000:>10.1f}"
            f"{total:>10.2f}{row['findings']:>8}"
            f"{row['bytes'] / 1024:>10.0f}"
        )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="DocAudit 性能基线（合成文档）")
    parser.add_argument(
        "--sizes",
        default=",".join(str(s) for s in DEFAULT_SIZES),
        help="页数列表，逗号分隔（默认 10,50,100）",
    )
    parser.add_argument(
        "--format",
        choices=["docx", "pptx", "both"],
        default="both",
        help="文档格式（默认 both）",
    )
    parser.add_argument("--seed", type=int, default=42, help="随机种子（默认 42，保证可复现）")
    parser.add_argument("--no-html", action="store_true", help="跳过 HTML 报告计时")
    parser.add_argument("--json", default=None, help="可选：把结果写入 JSON 文件")
    args = parser.parse_args(argv)

    try:
        sizes = [int(s) for s in str(args.sizes).split(",") if s.strip()]
    except ValueError:
        print(f"[ERROR] --sizes 需为逗号分隔整数: {args.sizes}")
        return 2
    formats = ("docx", "pptx") if args.format == "both" else (args.format,)

    print(f"[BENCH] sizes={sizes} formats={formats} seed={args.seed}")
    results = run_benchmark(
        sizes=sizes, formats=formats, seed=args.seed, with_html=not args.no_html
    )
    _print_table(results)

    if args.json:
        out = Path(args.json)
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"[BENCH] JSON 已写入: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
